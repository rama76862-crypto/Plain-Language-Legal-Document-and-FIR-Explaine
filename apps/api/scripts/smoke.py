"""Synthetic local HTTP journey. Never prints document text or session tokens."""

from __future__ import annotations

import argparse
import io
import time
from urllib.parse import urlparse

import httpx
import pymupdf
from docx import Document


JOURNEY_STEPS = 9


def synthetic(kind: str) -> tuple[bytes, str]:
    if kind == "pdf":
        doc = pymupdf.open()
        doc.new_page().insert_text(
            (50, 80), "SYNTHETIC agreement. Payment is due within 15 days of receipt."
        )
        data = doc.tobytes()
        doc.close()
        return data, "application/pdf"
    docx = Document()
    docx.add_paragraph("SYNTHETIC agreement. Payment is due within 15 days of receipt.")
    docx.add_table(rows=1, cols=1).cell(0, 0).text = "SYNTHETIC agreement table."
    stream = io.BytesIO()
    docx.save(stream)
    return stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _loopback(base: str) -> None:
    if urlparse(base).hostname not in {"localhost", "127.0.0.1", "::1"}:
        raise SystemExit("Smoke requires loopback services; use no public proxy or remote model.")


def _show_step(completed: int, label: str) -> None:
    width = 24
    filled = round(width * completed / JOURNEY_STEPS)
    bar = "#" * filled + "-" * (width - filled)
    print(f"[{bar}] {completed}/{JOURNEY_STEPS} {label}", flush=True)


def _show_clause_progress(completed: int, total: int, status: str) -> None:
    width = 24
    filled = round(width * completed / total) if total else 0
    bar = "#" * filled + "-" * (width - filled)
    print(f"  clauses [{bar}] {completed}/{total} status={status}", flush=True)


def run(base: str, private_base: str, kind: str, mode: str, extraction_only: bool = False) -> bool:
    _loopback(base)
    _loopback(private_base)
    with httpx.Client(base_url=base.rstrip("/"), timeout=180, trust_env=False) as client:
        health = client.get("/api/v1/health")
        health.raise_for_status()
        health_body = health.json()
        if health_body["mode"] != mode:
            raise SystemExit("Server mode differs from explicit smoke expectation.")
        if mode == "real" and not health_body["ml"].get("available") and not extraction_only:
            print("[------------------------] Real ML unavailable; no model result was fabricated.")
            return False
        if mode == "real" and not extraction_only:
            if health_body["ml"].get("classifier_ready") is not True:
                raise RuntimeError("BE2 classifier readiness was not confirmed.")
            if health_body["ml"].get("generator_ready") is not True:
                raise RuntimeError("BE2 generator readiness was not confirmed.")
        data, mime = synthetic(kind)
        response = client.post(
            "/api/v1/documents",
            files={"file": ("synthetic." + kind, data, mime)},
            data={"input_language": "en", "document_type": "agreement", "processing_consent": "true"},
        )
        response.raise_for_status()
        upload = response.json()
        headers = {"X-Session-Token": upload["session_token"]}
        path = "/api/v1/documents/" + upload["document_id"]
        _show_step(1, "loopback health confirms explicit service mode")
        _show_step(2, "synthetic upload accepted")
        analysis_id = None
        deleted = False
        try:
            deadline = time.monotonic() + 30
            while True:
                metadata_response = client.get(path, headers=headers)
                metadata_response.raise_for_status()
                metadata = metadata_response.json()
                if metadata["status"] == "failed":
                    raise RuntimeError(f"Extraction failed; safe metadata warnings: {metadata.get('warnings', [])}")
                if metadata["extraction_ready"]:
                    break
                if time.monotonic() > deadline:
                    raise RuntimeError("Extraction smoke timeout")
                time.sleep(0.1)
            extraction_response = client.get(path + "/extraction", headers=headers)
            extraction_response.raise_for_status()
            clauses = extraction_response.json()["clauses"]
            assert clauses
            if kind == "docx":
                assert all(clause["source_ref"]["page_number"] is None for clause in clauses)
            else:
                assert clauses[0]["source_ref"]["page_number"] == 1
            _show_step(3, f"native extraction ready; {len(clauses)} clause(s), source references checked")

            review = client.patch(
                path + "/extraction",
                headers=headers,
                json=dict(
                    expected_revision=upload["revision"],
                    corrections=[],
                    input_language="en",
                    document_type="agreement",
                ),
            )
            review.raise_for_status()
            revision = review.json()["revision"]
            assert review.json()["reviewed_revision"] == revision
            _show_step(4, "explicit unchanged-text review confirmed")

            source = client.get(path + "/source", headers=headers)
            source.raise_for_status()
            assert source.content == data
            _show_step(5, "authenticated original source verified")

            if extraction_only:
                _show_step(6, "analysis intentionally skipped by extraction-only mode")
                return True

            accepted = client.post(
                path + "/analyses",
                headers=headers,
                json=dict(expected_revision=revision, output_language="en"),
            )
            accepted.raise_for_status()
            analysis_id = accepted.json()["analysis_id"]
            print("  analysis accepted; polling actual clause progress", flush=True)
            deadline = time.monotonic() + 180
            health_latencies: list[float] = []
            while True:
                poll = client.get(path + "/analyses/" + analysis_id, headers=headers)
                poll.raise_for_status()
                snapshot = poll.json()
                progress = snapshot["progress"]
                _show_clause_progress(
                    progress["completed_clauses"],
                    progress["total_clauses"],
                    snapshot["status"],
                )
                if snapshot["status"] not in {"queued", "running"}:
                    break
                health_started = time.perf_counter()
                live_health = client.get("/api/v1/health")
                health_latencies.append(time.perf_counter() - health_started)
                live_health.raise_for_status()
                live_health_body = live_health.json()
                if live_health_body.get("status") != "ok" or live_health_body.get("ml", {}).get("available") is not True:
                    raise RuntimeError("Health lost real classifier or generator readiness during analysis.")
                if time.monotonic() > deadline:
                    raise RuntimeError("Analysis smoke timeout")
                time.sleep(0.5)
            if snapshot["status"] == "failed" or not snapshot.get("result"):
                raise RuntimeError("No usable local analysis result; smoke failed honestly.")
            if mode == "fixture":
                assert snapshot["status"] == "needs_review"
                assert snapshot["result"]["coverage"]["analysed"] == 0
                assert "DEVELOPMENT FIXTURE" in str(snapshot["warnings"])
            by_id = {clause["clause_id"]: clause["reviewed_text"] for clause in clauses}
            for item in snapshot["result"]["items"]:
                for evidence in item["evidence"]:
                    assert evidence["quote"] in by_id[evidence["clause_id"]]
            max_health_ms = round(max(health_latencies, default=0.0) * 1000.0, 1)
            _show_step(
                6,
                f"terminal {snapshot['status']} result and source evidence validated; "
                f"live health checks={len(health_latencies)}, max={max_health_ms} ms",
            )

            grounded_response = client.post(
                path + "/questions",
                headers=headers,
                json=dict(
                    analysis_id=analysis_id,
                    question="When is payment due?",
                    output_language="en",
                    client_request_id="smoke-grounded-once",
                ),
            )
            grounded_response.raise_for_status()
            grounded = grounded_response.json()
            print(
                "  grounded-answer flags: "
                f"review_status={grounded.get('review_status')} "
                f"limitations={grounded.get('limitations', [])} "
                f"evidence_count={len(grounded.get('evidence', []))}",
                flush=True,
            )
            assert grounded["answer"].strip()
            assert grounded["review_status"] == "needs_review"
            assert "insufficient_evidence" not in grounded["limitations"]
            assert grounded["evidence"]
            for evidence in grounded["evidence"]:
                assert evidence["quote"] in by_id[evidence["clause_id"]]
            _show_step(7, "real local question answer includes exact reviewed-source evidence; semantic review remains required")

            unsupported_response = client.post(
                path + "/questions",
                headers=headers,
                json=dict(
                    analysis_id=analysis_id,
                    question="What is the weather on Jupiter?",
                    output_language="en",
                    client_request_id="smoke-unsupported-once",
                ),
            )
            unsupported_response.raise_for_status()
            unsupported = unsupported_response.json()
            assert unsupported["review_status"] == "needs_review"
            assert "insufficient_evidence" in unsupported["limitations"]
            assert unsupported["evidence"] == []
            assert unsupported["reference_ids"] == []
            _show_step(8, "unsupported question returned explicit insufficient-evidence response")
        finally:
            deleted_response = client.delete(path, headers=headers)
            if deleted_response.status_code != 404:
                deleted_response.raise_for_status()
            deleted = True
        assert deleted
        assert client.get(path, headers=headers).status_code == 404

        if analysis_id:
            deadline = time.monotonic() + 45
            with httpx.Client(base_url=private_base.rstrip("/"), timeout=5, trust_env=False) as private:
                while True:
                    private_status = private.get("/internal/jobs/" + analysis_id)
                    if private_status.status_code == 404:
                        break
                    if private_status.status_code == 410:
                        body = private_status.json()
                        if body.get("error", {}).get("code") != "JOB_DELETED":
                            raise RuntimeError("Private service returned an unexpected deletion response.")
                        break
                    if private_status.status_code != 200 or time.monotonic() > deadline:
                        raise RuntimeError("Private deletion tombstone was not confirmed.")
                    time.sleep(0.25)
        _show_step(9, "public session deleted; public and private job tombstones confirmed")
        print("Real local service journey completed. Generated analysis remains needs_review until semantic review.")
        return True


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", default="http://127.0.0.1:8000")
    parser.add_argument("--private-url", default="http://127.0.0.1:8001")
    parser.add_argument("--kind", choices=["pdf", "docx"], default="pdf")
    parser.add_argument("--mode", choices=["real", "fixture"], default="real")
    parser.add_argument("--extraction-only", action="store_true")
    arguments = parser.parse_args()
    raise SystemExit(
        0
        if run(
            arguments.base_url,
            arguments.private_url,
            arguments.kind,
            arguments.mode,
            arguments.extraction_only,
        )
        else 2
    )
