"""Public acceptance checks: real parsers, synthetic docs and explicit ML doubles."""
import asyncio
import io
import threading
import time
from copy import deepcopy

import httpx
import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

from apps.api.config import Settings
from apps.api.main import create_app
from apps.api.ml_client import FixtureMLClient, RealMLClient
from apps.api.schemas import ApiError, Clause, SourceRef


def pdf_bytes(text='1. SYNTHETIC: Payment is due within 15 days of receipt.'):
    doc=pymupdf.open(); doc.new_page().insert_text((50,80),text)
    data=doc.tobytes(); doc.close(); return data


def docx_bytes():
    doc=Document(); doc.add_paragraph('1. SYNTHETIC first clause.')
    doc.add_table(rows=1,cols=1).cell(0,0).text='SYNTHETIC table clause.'
    doc.add_paragraph('2. SYNTHETIC last clause.')
    stream=io.BytesIO(); doc.save(stream); return stream.getvalue()


@pytest.fixture
def settings(tmp_path):
    return Settings(runtime_dir=tmp_path/'runtime',ml_mode='fixture',poll_interval_seconds=.005,cleanup_interval_seconds=.01)


@pytest.fixture
def client(settings):
    with TestClient(create_app(settings)) as client:
        yield client


def upload(client,data=None,kind='pdf'):
    mime='application/pdf' if kind=='pdf' else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    response=client.post('/api/v1/documents',files={'file':('synthetic.'+kind,pdf_bytes() if data is None else data,mime)},
        data={'input_language':'en','document_type':'agreement','processing_consent':'true'})
    assert response.status_code==202,response.text
    value=response.json(); return value,{'X-Session-Token':value['session_token']}


def wait_extraction(client,value,headers):
    url='/api/v1/documents/'+value['document_id']
    for _ in range(2000):
        metadata=client.get(url,headers=headers).json()
        if metadata.get('extraction_ready') or metadata.get('status')=='failed':
            return client.get(url+'/extraction',headers=headers)
        time.sleep(.005)
    pytest.fail('Extraction did not finish')


def confirm(client,value,headers,revision=1,corrections=None):
    return client.patch('/api/v1/documents/'+value['document_id']+'/extraction',headers=headers,
        json={'expected_revision':revision,'corrections':corrections or [],'input_language':'en','document_type':'agreement'})


def start(client,value,headers,revision=1,language='en'):
    return client.post('/api/v1/documents/'+value['document_id']+'/analyses',headers=headers,
        json={'expected_revision':revision,'output_language':language})


def wait_analysis(client,value,headers,aid):
    url=f"/api/v1/documents/{value['document_id']}/analyses/{aid}"
    for _ in range(300):
        response=client.get(url,headers=headers)
        if response.json().get('status') not in {'queued','running'}:
            return response
        time.sleep(.005)
    pytest.fail('Analysis did not finish')


def test_native_pdf_review_source_analysis_question_delete(client,caplog):
    value,headers=upload(client); base='/api/v1/documents/'+value['document_id']
    assert start(client,value,headers).status_code==409
    extraction=wait_extraction(client,value,headers)
    assert extraction.status_code==200
    clauses=extraction.json()['clauses']; assert clauses[0]['source_ref']['page_number']==1
    assert confirm(client,value,headers).json()['revision']==1
    assert client.get(base+'/source').status_code==404
    source=client.get(base+'/source',headers=headers)
    assert source.content.startswith(b'%PDF') and source.headers['cache-control']=='no-store'
    accepted=start(client,value,headers); assert accepted.status_code==202
    aid=accepted.json()['analysis_id']
    snapshot=wait_analysis(client,value,headers,aid).json()
    assert snapshot['status']=='needs_review'
    assert snapshot['result']['coverage']['analysed']==0
    assert 'DEVELOPMENT FIXTURE' in str(snapshot)
    body=dict(analysis_id=aid,question='What does this say?',output_language='en',client_request_id='same')
    answer=client.post(base+'/questions',headers=headers,json=body)
    assert answer.status_code==200
    assert client.post(base+'/questions',headers=headers,json=body).json()==answer.json()
    assert client.post(base+'/questions',headers=headers,json={**body,'question':'different'}).status_code==409
    assert client.delete(base,headers=headers).status_code==204
    assert client.get(base,headers=headers).status_code==404
    assert not list(client.app.state.settings.runtime_dir.glob('doc_*'))
    assert value['session_token'] not in caplog.text
    assert 'Payment is due' not in caplog.text


def test_docx_order_revision_and_invalidation(client):
    value,headers=upload(client,docx_bytes(),'docx')
    extraction=wait_extraction(client,value,headers).json(); clauses=extraction['clauses']
    assert all(c['source_ref']['page_number'] is None for c in clauses)
    texts=' '.join(c['source_text'] for c in clauses)
    assert texts.index('first')<texts.index('table')<texts.index('last')
    assert confirm(client,value,headers).status_code==200
    aid=start(client,value,headers,language='hi').json()['analysis_id']
    wait_analysis(client,value,headers,aid)
    response=confirm(client,value,headers,corrections=[{'clause_id':clauses[0]['clause_id'],'text':'SYNTHETIC corrected text.'}])
    assert response.json()['revision']==2
    assert confirm(client,value,headers).status_code==409
    base='/api/v1/documents/'+value['document_id']
    assert client.get(base+'/analyses/'+aid,headers=headers).status_code==404
    assert client.get(base+'/extraction',headers=headers).json()['clauses'][0]['source_ref']['bbox'] is None


def test_tokens_validation_and_upload_errors(client):
    a,ha=upload(client); b,hb=upload(client)
    for suffix in ['', '/source','/extraction','/analyses/nope']:
        assert client.get('/api/v1/documents/'+b['document_id']+suffix,headers=ha).status_code==404
    assert client.delete('/api/v1/documents/'+b['document_id'],headers=ha).status_code==404
    assert client.post('/api/v1/documents',files={'file':('x.pdf',b'','application/pdf')}).status_code==422
    assert client.post('/api/v1/documents',files={'file':('x.exe',b'x','application/octet-stream')}).status_code==415
    malformed,h=upload(client,b'not a pdf'); response=wait_extraction(client,malformed,h)
    assert response.status_code==422
    response=client.post('/api/v1/documents',files={'file':('x.pdf',pdf_bytes(),'application/pdf')},data={'input_language':'secret-invalid-value'})
    assert response.status_code==422 and response.json()['error']['code']=='VALIDATION_ERROR'
    assert 'secret-invalid-value' not in response.text


def test_oversize_and_health(settings):
    settings.max_upload_mb=1
    with TestClient(create_app(settings)) as client:
        r=client.post('/api/v1/documents',files={'file':('x.pdf',b'x'*(1024*1024+1),'application/pdf')},data={'processing_consent':'true'})
        assert r.status_code==413
        h=client.get('/api/v1/health').json()
        assert h['mode']=='fixture' and h['ml']['simulation'] and not h['ml']['available']
        assert h['extraction']['ready']


class ControlledML(FixtureMLClient):
    def __init__(self,settings):
        super().__init__(settings); self.deleted=[]; self.question_calls=0; self.fail_delete=0
    async def delete(self,aid):
        if self.fail_delete:
            self.fail_delete-=1
            raise ApiError('MODEL_UNAVAILABLE','Unavailable',503,True)
        self.deleted.append(aid); await super().delete(aid)
    async def question(self,payload):
        self.question_calls+=1
        await asyncio.sleep(.02)
        return await super().question(payload)


def test_all_language_jobs_cleanup_and_retry(settings):
    ml=ControlledML(settings)
    with TestClient(create_app(settings,ml_client=ml)) as client:
        value,headers=upload(client); wait_extraction(client,value,headers); confirm(client,value,headers)
        aids=[]
        for lang in ['en','hi','pa']:
            aid=start(client,value,headers,language=lang).json()['analysis_id']; aids.append(aid)
            wait_analysis(client,value,headers,aid)
        ml.fail_delete=1
        assert client.delete('/api/v1/documents/'+value['document_id'],headers=headers).status_code==204
        client.portal.call(client.app.state.cleanup_once)
        assert client.app.state.cleanup
        for entry in client.app.state.cleanup.values(): entry['next']=0
        client.portal.call(client.app.state.cleanup_once)
        assert set(aids).issubset(ml.deleted)
        assert not client.app.state.cleanup


def test_delete_during_extraction_no_resurrection(settings):
    began=threading.Event(); release=threading.Event()
    def slow(path,kind,language,settings,progress):
        began.set(); release.wait(2)
        return dict(clauses=[Clause(clause_id='c1',source_text='Synthetic',reviewed_text='Synthetic',source_ref=SourceRef(page_number=1))],warnings=[],input_language='en',document_type='agreement')
    with TestClient(create_app(settings,extractor=slow)) as client:
        value,headers=upload(client); assert began.wait(1)
        base='/api/v1/documents/'+value['document_id']
        assert client.delete(base,headers=headers).status_code==204
        release.set(); time.sleep(.03)
        assert client.get(base,headers=headers).status_code==404
        assert not list(settings.runtime_dir.glob('doc_*'))


def test_real_http_adapter_public_journey(settings):
    """Real HTTP adapter with MockTransport, not a real model proof."""
    settings.ml_mode='real'; private={}; question_count=0
    def handler(request):
        nonlocal question_count
        import json
        if request.url.path=='/internal/health':
            return httpx.Response(200,json={'classifier_loaded':True,'generator_loaded':True,'tested_language_matrix':[]})
        if request.method=='DELETE': return httpx.Response(204)
        if request.url.path=='/internal/analyse':
            payload=json.loads(request.content); private[payload['analysis_id']]=payload
            return httpx.Response(202,json=dict(analysis_id=payload['analysis_id'],status='queued'))
        if request.url.path.startswith('/internal/jobs/'):
            p=private[request.url.path.rsplit('/',1)[-1]]
            items=[dict(item_id=c['clause_id'],clause_id=c['clause_id'],kind='explanation',category=None,explanation='Mock HTTP test explanation',
                evidence=[dict(clause_id=c['clause_id'],quote=c['reviewed_text'])],reference_ids=[],review_status='supported',detector='test',model_score=None,limitations=[]) for c in p['clauses']]
            result={k:p[k] for k in ['document_id','analysis_id','document_revision','input_language','output_language','document_type']}
            result.update(coverage=dict(analysed=len(items),total=len(items)),items=items,deadlines=[],warnings=[],legal_aid=[],provenance=dict(classifier_version='mock',generator='mock-http-only',rules_version='mock'))
            return httpx.Response(200,json=dict(analysis_id=p['analysis_id'],document_revision=p['document_revision'],output_language=p['output_language'],status='complete',progress=dict(completed_clauses=len(items),total_clauses=len(items)),result=result,warnings=[]))
        if request.url.path=='/internal/question':
            question_count+=1; p=json.loads(request.content)
            assert 'session_token' not in p
            assert set(p)=={'document_id','analysis_id','document_revision','question','output_language','client_request_id'}
            c=private[p['analysis_id']]['clauses'][0]
            return httpx.Response(200,json=dict(answer='Mock grounded answer',evidence=[dict(clause_id=c['clause_id'],quote=c['reviewed_text'])],reference_ids=[],review_status='supported',limitations=['HTTP contract test only']))
        return httpx.Response(404)
    ml=RealMLClient(settings,transport=httpx.MockTransport(handler))
    with TestClient(create_app(settings,ml_client=ml)) as client:
        value,headers=upload(client); wait_extraction(client,value,headers); confirm(client,value,headers)
        aid=start(client,value,headers).json()['analysis_id']
        assert wait_analysis(client,value,headers,aid).json()['status']=='complete'
        body=dict(analysis_id=aid,question='Explain',output_language='en',client_request_id='retry')
        url='/api/v1/documents/'+value['document_id']+'/questions'
        assert client.post(url,headers=headers,json=body).status_code==200
        assert client.post(url,headers=headers,json=body).status_code==200
        assert question_count==1


def test_expiry_and_polling_does_not_extend(settings):
    with TestClient(create_app(settings)) as client:
        value,headers=upload(client); wait_extraction(client,value,headers)
        record=client.app.state.store.records[value['document_id']]
        before=record.expires_at
        client.get('/api/v1/documents/'+value['document_id'],headers=headers)
        assert record.expires_at==before
        record.expires_at=time.time()-1
        assert client.get('/api/v1/documents/'+value['document_id'],headers=headers).status_code==404
        time.sleep(.03)
        assert not list(settings.runtime_dir.glob('doc_*'))


def test_openapi_has_completion_operations():
    schema=create_app().openapi()
    assert '/api/v1/capabilities' in schema['paths']
    assert '/api/v1/documents/{document_id}/analyses/{analysis_id}/report' in schema['paths']
    assert 'delete' in schema['paths']['/api/v1/documents/{document_id}/analyses/{analysis_id}']
